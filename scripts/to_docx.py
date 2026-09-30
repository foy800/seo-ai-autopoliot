#!/usr/bin/env python3
"""Собирает Word-файл из статьи в Markdown: заголовки, жирные врезки, списки, таблицы, обложка.

  python3 to_docx.py out/2026-09-30-slug.md [--image out/2026-09-30-slug.png] [--out файл.docx]

Нужен пакет python-docx (pip install python-docx). Остальному скиллу он не нужен.
"""
import argparse
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from article_check import parse_article  # noqa: E402


def add_inline(par, text):
    """Жирный **текст** и ссылки [текст](url) внутри абзаца."""
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*|\[([^\]]+)\]\(([^)\s]+)\)", text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        if m.group(1) is not None:
            par.add_run(m.group(1)).bold = True
        else:
            par.add_run(f"{m.group(2)} ({m.group(3)})")
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def build(md_path, image, out):
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError:
        sys.exit("Нужен python-docx: pip install python-docx")
    raw = open(md_path, encoding="utf-8").read()
    meta, body = parse_article(raw)
    doc = Document()
    if image and os.path.exists(image) and os.path.getsize(image) > 100:
        try:
            doc.add_picture(image, width=Inches(6))
        except Exception:
            pass  # битую или служебную картинку пропускаем
    if meta.get("META_DESCRIPTION"):
        p = doc.add_paragraph()
        p.add_run("Description: ").bold = True
        p.add_run(meta["META_DESCRIPTION"])
    lines = body.splitlines()
    i = 0
    sep = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)+\|?\s*$")
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1
        elif ln.startswith("#"):
            level = len(ln) - len(ln.lstrip("#"))
            doc.add_heading(ln.lstrip("#").strip(), level=min(level, 3) if level > 1 else 1)
            i += 1
        elif "|" in ln and i + 1 < len(lines) and sep.match(lines[i + 1]):
            rows = [ln]
            j = i + 2
            while j < len(lines) and "|" in lines[j] and lines[j].strip():
                rows.append(lines[j])
                j += 1
            cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
            t = doc.add_table(rows=len(cells), cols=len(cells[0]))
            t.style = "Table Grid"
            for ri, row in enumerate(cells):
                for ci, val in enumerate(row[: len(cells[0])]):
                    cell = t.cell(ri, ci)
                    cell.text = ""
                    run = cell.paragraphs[0].add_run(val)
                    run.bold = ri == 0
            doc.add_paragraph()
            i = j
        elif re.match(r"^\s*[-*+]\s+", ln):
            add_inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*+]\s+", "", ln))
            i += 1
        elif re.match(r"^\s*\d+[.)]\s+", ln):
            add_inline(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", ln))
            i += 1
        else:
            par_lines = [ln]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\s*[-*+]\s|\s*\d+[.)]\s)", lines[i]) and "|" not in lines[i]:
                par_lines.append(lines[i])
                i += 1
            add_inline(doc.add_paragraph(), " ".join(x.strip() for x in par_lines))
    doc.save(out)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("md")
    ap.add_argument("--image", default=None)
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    image = a.image or re.sub(r"\.md$", ".png", a.md)
    out = a.out or re.sub(r"\.md$", ".docx", a.md)
    print(build(a.md, image, out))


if __name__ == "__main__":
    main()
