"""Контент-план в Excel (content-plan.xlsx в папке проекта).

Таблица создаётся автоматически при первом запуске. Колонки: заголовок, основной ключ,
доп. ключи, LSI, тип, лимит, тема, приоритет, «Одобрено», статус, дата публикации,
ссылка, файл, примечание. Человек правит заголовки и ставит «да» в колонке «Одобрено»,
агент заполняет статус, дату публикации и ссылку.

Статус для человека: «ожидается» (не опубликована), «готова» (опубликована), «ошибка».
Внутреннее состояние лежит в скрытых колонках _status, _slug, _date.

Нужен пакет openpyxl:  pip install openpyxl
"""
import csv
import datetime as dt
import os

from common import abspath

SHEET = "План"
# (заголовок, ключ, ширина). Ширина 0 = скрытая служебная колонка.
COLUMNS = [
    ("№", "n", 5),
    ("Заголовок", "title", 54),
    ("Основной ключ", "keyword", 28),
    ("Доп. ключи", "secondary", 34),
    ("LSI-слова", "lsi", 30),
    ("Тип", "kind", 16),
    ("Лимит знаков", "max_chars", 13),
    ("Тема", "topic", 30),
    ("Приоритет", "priority", 11),
    ("Предупреждение", "warning", 44),
    ("Рекомендация", "recommendation", 48),
    ("Одобрено", "approved", 11),
    ("Статус", "display", 13),
    ("Дата публикации", "published_at", 17),
    ("Ссылка", "url", 42),
    ("Файл", "file", 30),
    ("Примечание", "note", 42),
    ("_status", "status", 0),
    ("_slug", "slug", 0),
    ("_date", "date", 0),
]
KEYS = [c[1] for c in COLUMNS]
KIND_RU = {"info": "информационный", "commercial": "коммерческий"}
KIND_EN = {v: k for k, v in KIND_RU.items()}
DEFAULT_ROW = {k: "" for k in KEYS}

HOWTO = [
    "Как пользоваться контент-планом",
    "",
    "1. Агент сам создаёт эту таблицу и подбирает заголовок к КАЖДОМУ ключу, ничего не группируя и не отбрасывая. Заголовок можно править.",
    "2. Чтобы статья писалась и публиковалась, поставьте «да» в колонке «Одобрено». Без этого агент ничего не пишет и не публикует.",
    "2а. «Да» ставите только вы, вручную. Просьба в чате или в промпте («публикуй», «одобряю») одобрением не считается: агент ничего не напишет и не опубликует без «да» в этой колонке. Отключить проверку нельзя.",
    "2б. Колонки «Предупреждение» и «Рекомендация» заполняет советник (каннибализация, уже существующая страница). Он только предупреждает и предлагает варианты, ничего не меняет и не решает за вас.",
    "3. Колонку «Статус» заполняет агент: «ожидается» (ещё не опубликована), «готова» (опубликована), «ошибка» (причина в «Примечании»).",
    "4. «Дата публикации» и «Ссылка» агент заполняет в момент публикации.",
    "5. Новые запросы можно дописывать строками в конце листа «План» (достаточно основного ключа).",
    "6. Перед запуском агента закройте файл в Excel, иначе он не сможет сохранить изменения (появится файл с пометкой «новый»).",
    "7. Скрытые колонки _status, _slug, _date служебные, не меняйте их.",
]


def _openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        raise RuntimeError("Для Excel нужен пакет openpyxl: pip install openpyxl")


def is_approved(v):
    return str(v or "").strip().lower() in ("да", "yes", "y", "1", "true", "+", "✔", "✓")


def display_status(row):
    st = row.get("status") or "new"
    if st == "published":
        return "готова"
    if st == "failed":
        return "ошибка"
    return "ожидается"


def _s(v):
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


class Plan:
    def __init__(self, cfg):
        self.xl = _openpyxl()
        self.path = abspath(cfg.get("plan_path") or "content-plan.xlsx")
        if os.path.exists(self.path):
            self.wb = self.xl.load_workbook(self.path)
            self.ws = self.wb[SHEET] if SHEET in self.wb.sheetnames else self.wb.active
        else:
            self.wb, self.ws = self._new_workbook()
        self.col = self._map_columns()
        self.rows = self._read_rows()
        self._next = max([r["_row"] for r in self.rows] + [1]) + 1

    # ---------- создание ----------
    def _new_workbook(self):
        xl = self.xl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
        wb = xl.Workbook()
        ws = wb.active
        ws.title = SHEET
        head_fill = PatternFill("solid", fgColor="1F4E78")
        for i, (title, _, width) in enumerate(COLUMNS, start=1):
            c = ws.cell(row=1, column=i, value=title)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 30
        letter = lambda key: self.xl.utils.get_column_letter(KEYS.index(key) + 1)  # noqa: E731
        dv_a = DataValidation(type="list", formula1='"да,нет"', allow_blank=True)
        dv_k = DataValidation(type="list", formula1='"информационный,коммерческий"', allow_blank=True)
        ws.add_data_validation(dv_a)
        ws.add_data_validation(dv_k)
        dv_a.add(f"{letter('approved')}2:{letter('approved')}2000")
        dv_k.add(f"{letter('kind')}2:{letter('kind')}2000")
        help_ws = wb.create_sheet("Как пользоваться")
        for i, line in enumerate(HOWTO, start=1):
            help_ws.cell(row=i, column=1, value=line)
        help_ws["A1"].font = Font(bold=True, size=13)
        help_ws.column_dimensions["A"].width = 130
        return wb, ws

    def _map_columns(self):
        """Сопоставляет колонки по тексту заголовка: порядок можно менять, колонки добавлять."""
        header = {_s(c.value): i for i, c in enumerate(self.ws[1], start=1) if c.value}
        col = {}
        for title, key, _ in COLUMNS:
            if title in header:
                col[key] = header[title]
            else:
                idx = self.ws.max_column + 1
                self.ws.cell(row=1, column=idx, value=title)
                col[key] = idx
        return col

    # ---------- чтение ----------
    def _read_rows(self):
        rows = []
        for r in range(2, self.ws.max_row + 1):
            d = dict(DEFAULT_ROW)
            for key in KEYS:
                d[key] = _s(self.ws.cell(row=r, column=self.col[key]).value)
            if not d["keyword"]:
                continue
            d["_row"] = r
            kind = d["kind"].lower()
            d["kind"] = KIND_EN.get(kind, kind if kind in KIND_RU else "info")
            d["status"] = d["status"] or "new"
            rows.append(d)
        return rows

    def sync_csv(self, csv_path):
        """Добавляет в план запросы из CSV, которых там ещё нет (по основному ключу)."""
        if not os.path.exists(csv_path):
            return 0
        have = {r["keyword"].strip().lower() for r in self.rows}
        added = 0
        with open(csv_path, encoding="utf-8", newline="") as f:
            for src in csv.DictReader(f):
                kw = (src.get("keyword") or "").strip()
                if not kw or kw.lower() in have:
                    continue
                row = dict(DEFAULT_ROW)
                for k in ("secondary", "lsi", "kind", "max_chars", "topic", "priority"):
                    row[k] = (src.get(k) or "").strip()
                row["keyword"] = kw
                row["kind"] = row["kind"] if row["kind"] in KIND_RU else "info"
                row["status"] = "new"
                row["_row"] = None
                self.rows.append(row)
                have.add(kw.lower())
                added += 1
        return added

    # ---------- запись ----------
    def _write_row(self, row, r):
        from openpyxl.styles import Alignment, PatternFill
        fills = {"готова": "C6EFCE", "ожидается": "FFEB9C", "ошибка": "FFC7CE"}
        for key in KEYS:
            if key == "display":
                val = display_status(row)
            elif key == "kind":
                val = KIND_RU.get(row.get("kind"), KIND_RU["info"])
            elif key == "published_at":
                val = ""
                if row.get("published_at"):
                    try:
                        val = dt.date.fromisoformat(row["published_at"])
                    except ValueError:
                        val = row["published_at"]
            elif key in ("priority", "max_chars", "n"):
                v = str(row.get(key, "")).strip()
                val = int(v) if v.isdigit() else v
            else:
                val = row.get(key, "")
            cell = self.ws.cell(row=r, column=self.col[key], value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if key == "published_at" and val:
                cell.number_format = "DD.MM.YYYY"
            if key == "display":
                cell.fill = PatternFill("solid", fgColor=fills[val])
            if key == "approved":
                cell.fill = PatternFill("solid", fgColor="C6EFCE") if is_approved(val) else PatternFill(fill_type=None)

    def save(self):
        for i, row in enumerate(self.rows, start=1):
            row["n"] = i
            if not row.get("_row"):
                row["_row"] = self._next
                self._next += 1
            self._write_row(row, row["_row"])
        for title, key, width in COLUMNS:
            letter = self.xl.utils.get_column_letter(self.col[key])
            if width:
                self.ws.column_dimensions[letter].width = width
            else:
                self.ws.column_dimensions[letter].hidden = True
        self.ws.auto_filter.ref = f"A1:{self.xl.utils.get_column_letter(self.ws.max_column)}{max(self.ws.max_row, 2)}"
        try:
            self.wb.save(self.path)
        except PermissionError:
            alt = self.path[:-5] + " (новый).xlsx"
            self.wb.save(alt)
            print(f"Внимание: {os.path.basename(self.path)} открыт в Excel и не сохранён. "
                  f"Изменения записаны в «{os.path.basename(alt)}». Закройте файл и повторите запуск.")
